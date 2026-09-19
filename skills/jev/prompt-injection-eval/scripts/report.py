"""Write results as .xlsx, .csv and .json."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

DISCLAIMER = (
    "RESEARCH PROOF OF CONCEPT. The risk score is TypeSafe Jev's answer to a fixed set "
    "of questions about one prompt, weighted in code. It is a signal for a security "
    "gate, not a guarantee: no detector catches every injection, and this one should "
    "sit alongside least-privilege tools, output checks and logging, not replace them. "
    "Evaluate it on a labelled sample of your own traffic, tune the thresholds, and run "
    "it in alert mode (log and review, never block) until the false-alarm rate on real "
    "users is known. Only then consider block mode."
)

SIMPLE_COLUMNS = [
    ("Decision", "decision", 9),
    ("Risk score (Jev)", "risk", 14),
    ("Attack type (Jev)", "attack_type", 18),
    ("Signals (Jev)", "signals_text", 40),
    ("Prompt", "prompt", 90),
    ("Regex baseline (code)", "regex_decision", 18),
    ("LLM judge", "llm_verdict", 12),
    ("Label", "label", 10),
    ("Category", "category", 18),
]

DETAIL_COLUMNS = SIMPLE_COLUMNS[:5] + [
    ("Severity (0-3)", "manipulation_severity_raw", 12),
    ("Severity confidence", "manipulation_severity_confidence", 16),
    ("Override P", "instruction_override", 10),
    ("Role hijack P", "role_hijack", 12),
    ("Secret extraction P", "secret_extraction", 16),
    ("Action misuse P", "action_misuse", 14),
    ("False authority P", "false_authority", 15),
    ("Obfuscation P", "obfuscation", 13),
    ("Embedded instr. P", "embedded_instructions", 15),
    ("Attack type confidence", "attack_type_confidence", 18),
    ("Jev latency ms", "latency_ms", 13),
    ("Jev input tokens", "input_tokens", 14),
    ("Regex hits", "regex_hits_text", 40),
    ("LLM judge", "llm_verdict", 12),
    ("LLM latency ms", "llm_latency_ms", 13),
    ("LLM input tokens", "llm_input_tokens", 14),
    ("LLM output tokens", "llm_output_tokens", 15),
    ("Label", "label", 10),
    ("Category", "category", 18),
    ("Id", "id", 8),
    ("Error", "error", 40),
]

DECISION_FILL = {"block": "F8CBAD", "review": "FFE699", "allow": "C6E0B4"}


def write_all(out_dir: Path, rows: list[dict], summary: dict, comparison: list[dict]) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {"json": out_dir / "results.json", "csv": out_dir / "results.csv",
             "detail_csv": out_dir / "results-detail.csv"}
    paths["json"].write_text(json.dumps(
        {"summary": summary, "comparison": comparison, "disclaimer": DISCLAIMER,
         "results": rows}, indent=2))
    _write_csv(paths["csv"], rows, SIMPLE_COLUMNS)
    _write_csv(paths["detail_csv"], rows, DETAIL_COLUMNS)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("note: openpyxl not installed, skipping results.xlsx", file=sys.stderr)
        return paths
    paths["xlsx"] = out_dir / "results.xlsx"
    _write_xlsx(paths["xlsx"], rows, summary, comparison)
    return paths


def _write_csv(path: Path, rows: list[dict], columns: list[tuple]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(h for h, _, _ in columns)
        for row in rows:
            writer.writerow(row.get(k, "") for _, k, _ in columns)


def _write_xlsx(path: Path, rows: list[dict], summary: dict, comparison: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    bold = Font(bold=True)
    _results_sheet(book.active, "Results", rows, SIMPLE_COLUMNS, bold)
    _results_sheet(book.create_sheet("Details"), "Details", rows, DETAIL_COLUMNS, bold)
    _comparison_sheet(book, comparison, bold)
    _kv_sheet(book, "Summary", summary, bold)
    _text_sheet(book, "Read me", DISCLAIMER + "\n\n" + _column_guide())
    book.save(path)


def _results_sheet(sheet, title: str, rows: list[dict], columns: list[tuple], bold) -> None:
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.title = title
    for col, (header, _, width) in enumerate(columns, start=1):
        sheet.cell(row=1, column=col, value=header).font = bold
        sheet.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=2):
        for col, (_, key, _) in enumerate(columns, start=1):
            cell = sheet.cell(row=r, column=col, value=row.get(key, ""))
            if key in ("prompt", "signals_text", "regex_hits_text"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = DECISION_FILL.get(row.get("decision", ""))
        if fill:
            sheet.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def _comparison_sheet(book, comparison: list[dict], bold) -> None:
    if not comparison:
        return
    sheet = book.create_sheet("Comparison")
    headers = list(comparison[0].keys())
    for col, h in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=h).font = bold
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = 22
    for r, row in enumerate(comparison, start=2):
        for col, h in enumerate(headers, start=1):
            sheet.cell(row=r, column=col, value=row.get(h, ""))


def _kv_sheet(book, title: str, data: dict, bold) -> None:
    sheet = book.create_sheet(title)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 70
    for r, (k, v) in enumerate(data.items(), start=1):
        sheet.cell(row=r, column=1, value=k).font = bold
        sheet.cell(row=r, column=2, value=v if isinstance(v, (str, int, float, bool)) or v is None else json.dumps(v))


def _text_sheet(book, title: str, text: str) -> None:
    from openpyxl.styles import Alignment

    sheet = book.create_sheet(title)
    sheet.column_dimensions["A"].width = 120
    for r, line in enumerate(text.splitlines() or [""], start=1):
        sheet.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")


def _column_guide() -> str:
    return "\n".join([
        "Column guide",
        "Decision: allow / review / block from the risk score and thresholds in questions.py. Review also catches allows where Jev reported low confidence.",
        "Risk score (Jev): 0-100, Jev's answers to eight fixed questions, weighted in code.",
        "Attack type (Jev): Jev's pick from a fixed list, with 'none' for ordinary requests. Descriptive; the decision comes from the score.",
        "Signals (Jev): which yes/no questions came back at or above the signal threshold, strongest first.",
        "Regex baseline (code): what a phrase list would have done with the same prompt. Here for comparison only.",
        "LLM judge: what a generative model asked 'is this an injection?' replied, if that baseline was run. Here for comparison only.",
        "Label / Category: ground truth from the input file when provided; used for the Comparison sheet.",
        "Details sheet: every probability, confidence, latency and token count behind the Results sheet.",
    ])
