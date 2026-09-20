"""Write results as .xlsx, .csv and .json: one row per IP, attention first."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

DISCLAIMER = (
    "RESEARCH PROOF OF CONCEPT. Each row is TypeSafe Jev's answer to a fixed set of "
    "questions about a code-computed summary of one IP's requests, with rules in code "
    "raising the verdict where code is certain (exploit payloads on real routes, WAF "
    "denials). It is a triage signal, not an incident finding: a category of "
    "malicious means look at the requests, not block the address. Verify against "
    "the raw log before acting, keep a labelled sample of your own traffic to measure "
    "it on, and treat benign_user as 'nothing stood out', never as 'verified human'."
)

SIMPLE_COLUMNS = [
    ("Attention", "attention_text", 10),
    ("Category", "category", 16),
    ("Score (0-100)", "score", 12),
    ("Band", "band", 11),
    ("Confidence", "category_confidence", 11),
    ("IP", "ip", 18),
    ("Requests", "requests", 9),
    ("Window", "span_text", 22),
    ("Peak/min", "peak_rpm", 9),
    ("404 rate", "not_found_rate", 9),
    ("WAF deny", "waf_deny", 9),
    ("WAF labels", "waf_labels_text", 40),
    ("Signals (Jev)", "signals_text", 44),
    ("Code signals", "code_signals_text", 56),
    ("Top UA", "top_ua", 50),
    ("ASN", "asn", 8),
    ("Hosts", "hosts_text", 30),
    ("Sample paths", "paths_text", 70),
    ("Raw rows", "raw_file", 30),
    ("Label", "label", 16),
]

DETAIL_COLUMNS = SIMPLE_COLUMNS[:6] + [
    ("Jev category (raw)", "jev_category", 16),
    ("Rule reasons", "rule_reasons_text", 40),
    ("Jev severity level", "jev_severity_level", 16),
    ("Jev severity distribution", "jev_severity_distribution", 50),
    ("Jev severity expectation (0-3)", "jev_severity_expectation", 18),
    ("Jev severity confidence", "jev_severity_confidence", 16),
    ("Generic probing P", "generic_probing", 14),
    ("App aware P", "app_aware", 12),
    ("Exploit payloads P", "exploit_payloads", 15),
    ("Credential attack P", "credential_attack", 16),
    ("Enumeration P", "enumeration", 13),
    ("Automated P", "automated", 12),
    ("Declared bot P", "declared_bot", 13),
    ("AI operated P", "ai_operated", 13),
    ("Monitoring P", "monitoring", 12),
    ("Wrong host P", "wrong_host", 12),
    ("Scanner tool P", "scanner_tool", 13),
    ("Category probabilities", "probabilities_text", 50),
    ("Jev input tokens", "input_tokens", 14),
    ("Estimated tokens", "estimated_tokens", 14),
    ("Jev latency ms", "latency_ms", 13),
    ("Label", "label", 16),
    ("Error", "error", 40),
]

BAND_FILL = {"Attack": "F8CBAD", "Concerning": "FFD966", "Nuisance": "FFF2CC", "Benign": "E2EFDA"}
CATEGORY_FILL = {"malicious": "F8CBAD", "background_scan": "FFE699", "ai_agent": "BDD7EE",
                 "benign_bot": "DDEBF7", "benign_user": "C6E0B4", "unclear": "E7E6E6"}


def write_all(out_dir: Path, rows: list[dict], summary: dict) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {"json": out_dir / "results.json", "csv": out_dir / "results.csv",
             "detail_csv": out_dir / "results-detail.csv"}
    paths["json"].write_text(json.dumps({"summary": summary, "disclaimer": DISCLAIMER, "results": rows}, indent=2))
    _write_csv(paths["csv"], rows, SIMPLE_COLUMNS)
    _write_csv(paths["detail_csv"], rows, DETAIL_COLUMNS)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("note: openpyxl not installed, skipping results.xlsx", file=sys.stderr)
        return paths
    paths["xlsx"] = out_dir / "results.xlsx"
    _write_xlsx(paths["xlsx"], rows, summary)
    return paths


def _write_csv(path: Path, rows: list[dict], columns: list[tuple]) -> None:
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(h for h, _, _ in columns)
        for row in rows:
            writer.writerow(row.get(k, "") for _, k, _ in columns)


def _write_xlsx(path: Path, rows: list[dict], summary: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    bold = Font(bold=True)
    _results_sheet(book.active, "Results", rows, SIMPLE_COLUMNS, bold)
    _results_sheet(book.create_sheet("Details"), "Details", rows, DETAIL_COLUMNS, bold)
    _kv_sheet(book, "Summary", summary, bold)
    _text_sheet(book, "Read me", DISCLAIMER + "\n\n" + _column_guide())
    book.save(path)


def _results_sheet(sheet, title: str, rows: list[dict], columns: list[tuple], bold) -> None:
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.title = title
    for col, (header, _, width) in enumerate(columns, start=1):
        sheet.cell(row=1, column=col, value=header).font = bold
        sheet.column_dimensions[get_column_letter(col)].width = width
    wrap = {"signals_text", "code_signals_text", "paths_text", "top_ua", "hosts_text", "rule_reasons_text", "probabilities_text"}
    for r, row in enumerate(rows, start=2):
        for col, (_, key, _) in enumerate(columns, start=1):
            cell = sheet.cell(row=r, column=col, value=row.get(key, ""))
            if key in wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = CATEGORY_FILL.get(row.get("category", ""))
        if fill:
            sheet.cell(row=r, column=2).fill = PatternFill("solid", fgColor=fill)
        band_fill = BAND_FILL.get(row.get("band", ""))
        if band_fill and title == "Results":
            sheet.cell(row=r, column=3).fill = PatternFill("solid", fgColor=band_fill)
            sheet.cell(row=r, column=4).fill = PatternFill("solid", fgColor=band_fill)
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions


def _kv_sheet(book, title: str, data: dict, bold) -> None:
    sheet = book.create_sheet(title)
    r = 1
    for key, value in data.items():
        sheet.cell(row=r, column=1, value=key).font = bold
        if isinstance(value, dict):
            for k2, v2 in value.items():
                sheet.cell(row=r, column=2, value=str(k2))
                sheet.cell(row=r, column=3, value=v2 if isinstance(v2, (int, float)) else str(v2))
                r += 1
        else:
            sheet.cell(row=r, column=2, value=value if isinstance(value, (int, float)) else str(value))
            r += 1
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 60


def _text_sheet(book, title: str, text: str) -> None:
    from openpyxl.styles import Alignment
    sheet = book.create_sheet(title)
    for r, para in enumerate(text.split("\n"), start=1):
        cell = sheet.cell(row=r, column=1, value=para)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 120


def _column_guide() -> str:
    return (
        "Attention: YES when the score is 50 or more (Concerning band or above) and the category is not a benign one, or a code rule fired "
        "(exploit payloads or WAF attack signatures on routes that exist here; credential-attack volume on real auth endpoints). Start here.\n"
        "Category: benign_user | benign_bot | ai_agent | background_scan | malicious | unclear. Jev's choice, "
        "raised by code rules where code is certain (see Details > Rule reasons).\n"
        "Score (0-100): threat score computed in code from Jev's answers: 45% Jev's severity rating, 35% the strongest attack vector "
        "(exploit payloads, credential attack or enumeration) scaled by how much the traffic knows this application, 10% that app "
        "knowledge itself, 10% scanning pressure. The weights are in scripts/questions.py.\n"
        "Band: 0-24 Benign (ordinary use, declared bots, monitors), 25-49 Nuisance (scanning for software or files this site does not "
        "have), 50-74 Concerning (recon of real endpoints, sign-in attempts, WAF denials on real routes, enumeration), 75-100 Attack "
        "(exploit payloads against real endpoints, credential attacks at volume, enumeration returning successes). Bands derive from the "
        "score. The score is about threat; the category is about who. When they disagree, the Details sheet shows why.\n"
        "Signals (Jev): yes/no questions answered at 0.5 or above, strongest first.\n"
        "Code signals: facts regex and counting established before Jev was asked (scanner UA, probe families, payloads, WAF denials).\n"
        "Sample paths: the most-requested paths with their usual status, so you can judge the row without opening the log.\n"
        "Details sheet: every probability, the category distribution, tokens and latency per IP."
    )
